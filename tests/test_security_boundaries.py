"""Regression tests for upload confinement and spreadsheet text boundaries."""
from __future__ import annotations

import ast
import io
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
import xml.etree.ElementTree as ET
import zipfile

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import parser as catalog


def catalog_pdf(description="TOY COROLLA 2007"):
    """One real text PDF row; no external program or exploit payload."""
    text = f"VK-TEST {description} 100.00"
    stream = f"BT /F1 12 Tf 40 750 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(data)


def upload_function(temp_parent, parse=catalog.parse_catalog_pdf):
    # Importing the UI module starts Streamlit widgets and creates its SQLite DB.
    # Compile only the actual upload function so filesystem/parser behavior stays real.
    module = ast.parse((ROOT / "streamlit_app.py").read_text())
    function = next(node for node in module.body if isinstance(node, ast.FunctionDef) and node.name == "_cached_parse_catalog")
    function.decorator_list = []
    namespace = {
        "Path": Path,
        "tempfile": SimpleNamespace(TemporaryDirectory=lambda: tempfile.TemporaryDirectory(dir=temp_parent)),
        "parse_catalog_pdf": parse,
        "ParsedCatalog": catalog.ParsedCatalog,
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(ROOT / "streamlit_app.py"), "exec"), namespace)
    return namespace["_cached_parse_catalog"]


class UploadBoundaryTests(unittest.TestCase):
    def test_uploaded_paths_cannot_overwrite_outside_temporary_directory(self):
        # Removing basename confinement must make the outside sentinel change.
        for style in ("absolute", "traversal"):
            with self.subTest(style=style), tempfile.TemporaryDirectory() as outer:
                sentinel = Path(outer) / "outside.pdf"
                sentinel.write_bytes(b"untouched")
                name = str(sentinel) if style == "absolute" else "../outside.pdf"
                parsed = upload_function(outer)(name, catalog_pdf(), (), ())
                self.assertEqual(sentinel.read_bytes(), b"untouched")
                self.assertEqual(parsed.raw_rows[0].code, "VK-TEST")
                self.assertEqual(list(Path(outer).iterdir()), [sentinel])

    def test_normal_filename_preserves_brand_and_cleans_temporary_file(self):
        with tempfile.TemporaryDirectory() as outer:
            parsed = upload_function(outer)("NUVO catalog.pdf", catalog_pdf(), (), ())
            self.assertEqual(parsed.raw_rows[0].supplier_brand, "NUVO-PRO")
            self.assertEqual(parsed.raw_rows[0].price, 100.0)
            self.assertEqual(list(Path(outer).iterdir()), [])

    def test_parser_failure_cleans_temporary_file(self):
        def reject(path, **kwargs):
            self.assertTrue(Path(path).exists())
            raise ValueError("synthetic parse failure")
        with tempfile.TemporaryDirectory() as outer:
            with self.assertRaisesRegex(ValueError, "synthetic parse failure"):
                upload_function(outer, reject)("catalog.pdf", b"invalid", (), ())
            self.assertEqual(list(Path(outer).iterdir()), [])


class WorkbookBoundaryTests(unittest.TestCase):
    def test_supplier_pdf_description_round_trips_as_literal_text(self):
        # Direct assignment to openpyxl without string typing turns this into a formula.
        with tempfile.TemporaryDirectory() as outer:
            path = Path(outer) / "catalog.pdf"
            path.write_bytes(catalog_pdf("=1+1"))
            parsed = catalog.parse_catalog_pdf(path)
        self.assertEqual(parsed.raw_rows[0].description, "=1+1")
        raw = catalog.raw_rows_dataframe(parsed.raw_rows)
        workbook = catalog.build_workbook(pd.DataFrame(), raw)
        saved = load_workbook(io.BytesIO(catalog.workbook_to_bytes(workbook)), data_only=False)
        self.assertEqual(saved["Catalog Rows"]["F2"].value, "=1+1")
        self.assertEqual(saved["Catalog Rows"]["F2"].data_type, "s")

    def test_untrusted_text_on_every_sheet_preserves_only_generated_formulas(self):
        applications = pd.DataFrame([{
            "Model": "=1+1", "Original Price (PHP)": 100.0,
            "Your Price (PHP)": 0.0, "Use Price (PHP)": "=2+2",
            "Confidence": 0.95, "Plain": "+ordinary", "Missing": None,
        }])
        raw = pd.DataFrame([{"Description": '=HYPERLINK("https://example.invalid","catalog")'}])
        conflicts = pd.DataFrame([{"Issue": "=3+3"}])
        evidence = pd.DataFrame([{"Value": "=4+4"}])
        workbook = catalog.build_workbook(applications, raw, applications, conflicts, evidence)
        data = catalog.workbook_to_bytes(workbook)
        saved = load_workbook(io.BytesIO(data), data_only=False)
        self.assertEqual(saved.sheetnames, ["Applications", "Catalog Rows", "Review Queue", "Conflict Log", "Evidence Log"])
        for sheet, coordinate, value in [
            ("Applications", "A2", "=1+1"), ("Catalog Rows", "A2", '=HYPERLINK("https://example.invalid","catalog")'),
            ("Review Queue", "A2", "=1+1"), ("Conflict Log", "A2", "=3+3"), ("Evidence Log", "A2", "=4+4"),
        ]:
            with self.subTest(sheet=sheet):
                self.assertEqual(saved[sheet][coordinate].value, value)
                self.assertEqual(saved[sheet][coordinate].data_type, "s")
        for sheet in ("Applications", "Review Queue"):
            self.assertEqual(saved[sheet]["D2"].value, '=IF(C2="",B2,C2)')
            self.assertEqual(saved[sheet]["D2"].data_type, "f")
        self.assertEqual(saved["Applications"]["B2"].value, 100)
        self.assertEqual(saved["Applications"]["C2"].value, 0)
        self.assertIsNone(saved["Applications"]["G2"].value)
        self.assertEqual(saved["Applications"]["F2"].value, "+ordinary")
        self.assertEqual(saved["Applications"].freeze_panes, "A2")
        self.assertEqual(saved["Applications"].auto_filter.ref, "A1:G2")
        self.assertEqual(saved["Applications"]["E2"].number_format, "0%")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            formulas = []
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    formulas += [element.text for element in ET.fromstring(archive.read(name)).iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f")]
        self.assertEqual(formulas, ['IF(C2="",B2,C2)', 'IF(C2="",B2,C2)'])


if __name__ == "__main__":
    unittest.main()
